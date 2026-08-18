import collections
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    title_shape.text = title
    subtitle_shape.text = subtitle

def add_section_header(prs, title):
    slide_layout = prs.slide_layouts[2] # Section Header
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title

def add_content_slide(prs, title, content_points):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    
    for i, point in enumerate(content_points):
        if i == 0:
            p = tf.paragraphs[0]
            p.text = point
        else:
            p = tf.add_paragraph()
            p.text = point
            
    # Add a placeholder for screenshot
    left = Inches(5)
    top = Inches(2)
    width = Inches(4.5)
    height = Inches(3.5)
    
    shape = slide.shapes.add_shape(
        1, left, top, width, height # 1 is AutoShapeType.RECTANGLE
    )
    shape.text = "[PLACEHOLDER FOR SCREENSHOT]"
    # Make placeholder styling distinct
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(230, 230, 230)
    line = shape.line
    line.color.rgb = RGBColor(100, 100, 100)

prs = Presentation()

add_title_slide(prs, "CRM Application", "Complete User Manual for Admins and Agents")

# --- TENANT ADMIN SECTION ---
add_section_header(prs, "Part 1: Tenant Admin Manual")

add_content_slide(prs, "1. Admin Dashboard", [
    "High-level overview of total revenue and deals.",
    "Track team performance and agent activities.",
    "Identify bottlenecks in pipelines.",
    "Monitor recent activities across the organization."
])

add_content_slide(prs, "2. Lead Management", [
    "Add new leads manually or via import.",
    "Deduplication checks to prevent double entries.",
    "Edit lead details and assign to specific agents.",
    "Convert qualified leads into Accounts and Contacts."
])

add_content_slide(prs, "3. Contact & Account Management", [
    "360-degree view of Contact Profiles.",
    "Manage linked Deals, Activities, and Notes.",
    "View and edit Account (Company) details.",
    "Track all communications linked to an Account."
])

add_content_slide(prs, "4. Deals & Custom Pipelines", [
    "Kanban-style Pipeline view (Drag & Drop).",
    "Add and Edit deals, configure Deal Values.",
    "Move deals through custom stages.",
    "Mark Deals as Won or Lost (with Lost Reason tracking)."
])

add_content_slide(prs, "5. Omnichannel: Email & Templates", [
    "Create and manage Email Templates.",
    "Send automated or manual emails directly from CRM.",
    "Gmail integration to sync inbox and sent items.",
    "Configure Automated Sequences for drip campaigns."
])

add_content_slide(prs, "6. Omnichannel: Voice Calls", [
    "Built-in Quick Dialer using Twilio integration.",
    "View comprehensive Call Logs.",
    "Call statistics and metrics tracking.",
    "Click-to-call functionality from Contact profiles."
])

add_content_slide(prs, "7. Omnichannel: WhatsApp & Internal Chat", [
    "WhatsApp Bubbles for direct customer messaging.",
    "Full history of WhatsApp conversations.",
    "Teammates Chat for internal agent-to-agent collaboration.",
    "Internal Notes on Leads/Deals invisible to customers."
])

add_content_slide(prs, "8. Document Management", [
    "Upload contracts, invoices, and proposals.",
    "Preview documents directly in the browser.",
    "Link documents to specific Deals or Contacts."
])

add_content_slide(prs, "9. Productivity & Calendars", [
    "Google Calendar Sync for two-way event tracking.",
    "Create Tasks, Reminders, and Meetings.",
    "Receive Push Notifications for upcoming activities."
])

add_content_slide(prs, "10. Settings & User Management", [
    "Invite new Users and manage Roles (Admin vs. Agent).",
    "Detailed Permission Editor for granular access control.",
    "Configure custom Pipeline Stages.",
    "Manage app integrations (Google, Twilio, WhatsApp)."
])

add_content_slide(prs, "11. Privacy, Compliance & Support", [
    "Audit Logs to track all user actions.",
    "Consent Logs and Data Breach tracking for GDPR.",
    "Submit Support Tickets to the Platform Provider."
])

# --- AGENT SECTION ---
add_section_header(prs, "Part 2: Agent Manual")

add_content_slide(prs, "1. Agent Dashboard", [
    "Personalized view of assigned Tasks and Deals.",
    "Track individual performance and quotas.",
    "Quick access to today's scheduled activities."
])

add_content_slide(prs, "2. Managing Assigned Leads & Deals", [
    "View only Leads and Deals assigned to you.",
    "Update Deal stages as you progress.",
    "Add Internal Notes for context.",
    "Log manual activities (calls made, emails sent)."
])

add_content_slide(prs, "3. Communications (Agent)", [
    "Use Quick Dialer to call assigned Contacts.",
    "Send WhatsApp messages directly to Leads.",
    "Email Contacts using pre-approved Templates.",
    "Chat with Admins or peers via Teammates Chat."
])

add_content_slide(prs, "4. Calendar & Tasks", [
    "Manage your personal calendar.",
    "Check off completed tasks.",
    "Sync personal Google Calendar."
])

add_content_slide(prs, "5. Profile Settings", [
    "Update your personal profile information.",
    "Change your password securely."
])

prs.save("CRM_User_Manual.pptx")
print("Successfully generated CRM_User_Manual.pptx")
